#!/usr/bin/env python3
"""
================================================================================
MRPL DEEP KNOWLEDGE BASE CLASSIFIER
================================================================================

ANALYSIS-ONLY / READ-ONLY SCRIPT.

This script analyzes every file inside:

    D:\\MRPL-Sovereign-AI\\merged_knowledge_base\\08_Other

and produces three report files suggesting which of the following categories
each file most likely belongs to:

    01_Refinery_Manufacturing
    02_Environment_Compliance
    03_Safety_HSE
    04_Company_General
    05_Policies_Certifications
    06_Website_Content
    07_Finance
    08_Other

It NEVER deletes, moves, renames, or modifies any source file. It only
CREATES report files in the parent folder (merged_knowledge_base).

Reports produced:
    deep_classification_report.json
    deep_classification_report.csv
    deep_classification_summary.txt

Run with:
    python deep_classify_other.py
================================================================================
"""

import os
import sys
import io
import csv
import json
import hashlib
import traceback
import re
from pathlib import Path
from datetime import datetime
from collections import defaultdict, Counter
from difflib import SequenceMatcher

# ==============================================================================
# CONFIGURATION
# ==============================================================================

BASE_DIR = Path(r"D:\MRPL-Sovereign-AI\merged_knowledge_base")
SOURCE_DIR = BASE_DIR / "08_Other"

JSON_REPORT_PATH = BASE_DIR / "deep_classification_report.json"
CSV_REPORT_PATH = BASE_DIR / "deep_classification_report.csv"
SUMMARY_REPORT_PATH = BASE_DIR / "deep_classification_summary.txt"

CATEGORIES = [
    "01_Refinery_Manufacturing",
    "02_Environment_Compliance",
    "03_Safety_HSE",
    "04_Company_General",
    "05_Policies_Certifications",
    "06_Website_Content",
    "07_Finance",
]

ALL_CATEGORIES = CATEGORIES + ["08_Other"]

SUPPORTED_EXTENSIONS = {
    ".pdf", ".txt", ".docx", ".xlsx", ".xls", ".pptx", ".ppt",
    ".csv", ".json", ".html", ".htm",
}

# How many characters of "body" text to keep in memory per document (cap for
# very large documents). This is representative sampling, not the full doc.
MAX_BODY_CHARS = 60000
# How many pages of a PDF to treat as "first pages" for stronger weighting.
FIRST_PAGES_COUNT = 3
# Max pages to pull text from at all (representative sampling for huge PDFs).
MAX_PAGES_TO_SCAN = 40
# Max pages to attempt OCR on (OCR is expensive).
MAX_OCR_PAGES = 5
# OCR is only attempted if a PDF looks like it might be scanned.
MIN_CHARS_PER_PAGE_NOT_SCANNED = 40

GENERIC_WORDS = {
    "company", "project", "report", "policy", "management", "system",
    "document", "information", "general", "overview", "annual",
}
GENERIC_WORD_CAP = 2  # max score contribution from a generic word via frequency

BODY_FREQUENCY_CAP_PER_CATEGORY = 20

# ==============================================================================
# DEPENDENCY CHECKS
# ==============================================================================

DEPS = {
    "pymupdf": False,
    "pytesseract": False,
    "PIL": False,
    "docx": False,
    "openpyxl": False,
    "pptx": False,
}

try:
    import pymupdf  # noqa: F401  (NOT the deprecated `fitz` alias)
    DEPS["pymupdf"] = True
except ImportError:
    pass

try:
    import pytesseract  # noqa: F401
    DEPS["pytesseract"] = True
except ImportError:
    pass

try:
    from PIL import Image  # noqa: F401
    DEPS["PIL"] = True
except ImportError:
    pass

try:
    import docx  # python-docx  # noqa: F401
    DEPS["docx"] = True
except ImportError:
    pass

try:
    import openpyxl  # noqa: F401
    DEPS["openpyxl"] = True
except ImportError:
    pass

try:
    from pptx import Presentation  # python-pptx  # noqa: F401
    DEPS["pptx"] = True
except ImportError:
    pass

OCR_AVAILABLE = DEPS["pymupdf"] and DEPS["pytesseract"] and DEPS["PIL"]


def print_dependency_report():
    print("Dependency check:")
    for name, available in DEPS.items():
        status = "OK" if available else "MISSING"
        print(f"  - {name:12s}: {status}")
    if not OCR_AVAILABLE:
        print("  NOTE: OCR_NOT_AVAILABLE (pymupdf + pytesseract + Pillow all required for OCR)")
    print()


# ==============================================================================
# KEYWORD GROUPS
#
# Each category has STRONG signals (specific multi-word phrases / strong
# single terms unlikely to appear outside the domain) and MODERATE signals
# (weaker / more generic single terms). All keywords are lowercase.
# ==============================================================================

KEYWORD_GROUPS = {
    "01_Refinery_Manufacturing": {
        "strong": [
            "refinery", "refining", "process unit", "process plant",
            "plant operation", "cdu", "vdu", "hydrocracker",
            "hydrodesulfurization", "gohds", "polypropylene", "ppu",
            "petrochemical", "crude processing", "distillation",
            "process technology", "manufacturing unit", "refinery complex",
            "petroleum products", "naphtha", "kerosene", "atf",
            "aviation fuel", "bitumen", "hydrogen generation unit",
            "process equipment", "turnaround", "as 9100d",
            "crude oil processing", "product production",
        ],
        "moderate": [
            "manufacturing", "production", "polymer", "production capacity",
            "diesel", "gasoline", "petrol", "lpg", "sulfur", "utilities",
            "steam", "engineering", "plant",
        ],
    },
    "02_Environment_Compliance": {
        "strong": [
            "environment clearance", "environmental clearance",
            "ec compliance", "half yearly compliance",
            "environment statement", "moef", "moefcc",
            "pollution control", "kspcb", "consent to operate",
            "consent to establish", "hazardous waste",
            "waste authorization", "crz", "coastal regulation zone",
            "single point mooring", "environment monitoring",
            "environmental impact assessment", "eia report",
            "environment management plan", "environmental compliance",
        ],
        "moderate": [
            "environment protection", "cfo", "cto", "desalination", "spm",
            "apmc", "emission", "effluent", "wastewater", "air quality",
            "water quality", "environmental impact", "eia", "emp",
            "ecological", "compliance report",
        ],
    },
    "03_Safety_HSE": {
        "strong": [
            "health and safety", "occupational safety", "industrial safety",
            "fire safety", "emergency response", "permit to work",
            "confined space", "fire protection", "safety management",
            "process safety", "hazop", "lopa", "safety procedure",
            "emergency preparedness", "occupational health",
        ],
        "moderate": [
            "hse", "incident", "accident", "near miss", "hazard",
            "risk assessment", "ppe", "vigilance",
        ],
    },
    "04_Company_General": {
        "strong": [
            "mrpl overview", "corporate overview", "company profile",
            "board of directors", "company history", "about mrpl",
            "corporate structure", "organization structure",
            "corporate information",
        ],
        "moderate": [
            "vision", "mission", "organization", "management", "facilities",
            "capacity", "subsidiary", "hr", "human resources", "careers",
            "employee",
        ],
    },
    "05_Policies_Certifications": {
        "strong": [
            "information security policy", "record retention",
            "management system", "iso 9001", "iso 14001", "iso 45001",
            "iso 50001", "iso 27001", "iso/iec 17025", "certification",
            "accreditation", "quality management",
            "environmental management system", "energy management",
            "information security", "integrated management system",
            "compliance policy",
        ],
        "moderate": [
            "policy", "iso", "certificate", "audit",
        ],
    },
    "06_Website_Content": {
        "strong": [
            "official website of mrpl", "mrpl website", "website content",
            "html-derived content",
        ],
        "moderate": [
            "official website", "home", "about us", "web page",
            "navigation", "website-generated text", "page title",
        ],
    },
    "07_Finance": {
        "strong": [
            "annual report", "financial results", "financial statement",
            "balance sheet", "profit and loss", "cash flow statement",
            "annual general meeting", "postal ballot",
            "financial performance", "investor relations",
            "stock exchange", "corporate governance",
            "capital expenditure", "statutory auditor", "audit report",
            "financial year 2024-25", "financial year 2025-26",
            "fy 2024-25", "fy 2025-26",
        ],
        "moderate": [
            "revenue", "ebitda", "profit", "loss", "dividend",
            "shareholder", "agm", "board meeting", "investor", "bse",
            "nse", "sebi", "debt", "borrowing", "accounts",
            "financial year",
        ],
    },
}

SECTION_WEIGHTS = {
    "title": {"strong": 10, "moderate": 5},
    "filename": {"strong": 4, "moderate": 2},
    "first_page": {"strong": 8, "moderate": 4},
    "headings": {"strong": 7, "moderate": 3},
}
EXACT_PHRASE_BONUS = 8  # additional bonus for multi-word strong phrase found verbatim in body


# ==============================================================================
# TEXT EXTRACTION HELPERS
# ==============================================================================

def sha256_file(path, block_size=65536):
    h = hashlib.sha256()
    try:
        with open(path, "rb") as f:
            while True:
                block = f.read(block_size)
                if not block:
                    break
                h.update(block)
        return h.hexdigest()
    except Exception:
        return None


def clean_text(text):
    if not text:
        return ""
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


def guess_headings_from_text(text, max_headings=25):
    """
    Heuristic heading detection: short lines, mostly title-case or upper-case,
    not ending in punctuation typical of body sentences.
    """
    headings = []
    if not text:
        return headings
    lines = text.splitlines()
    for line in lines:
        line = line.strip()
        if not line:
            continue
        if len(line) < 3 or len(line) > 90:
            continue
        if line.endswith((".", ",", ";")):
            continue
        letters = [c for c in line if c.isalpha()]
        if not letters:
            continue
        upper_ratio = sum(1 for c in letters if c.isupper()) / len(letters)
        word_count = len(line.split())
        is_titleish = line.istitle() or upper_ratio > 0.6
        if is_titleish and word_count <= 12:
            headings.append(line)
        if len(headings) >= max_headings:
            break
    return headings


def extract_pdf(path):
    """
    Returns dict with: page_count, text_chars, text_words, first_page_text,
    body_text, headings, title_detected, document_type, ocr_attempted,
    ocr_success, notes
    """
    result = {
        "page_count": 0,
        "text_chars": 0,
        "text_words": 0,
        "first_page_text": "",
        "body_text": "",
        "headings": [],
        "title_detected": "",
        "document_type": "TEXT_PDF",
        "ocr_attempted": False,
        "ocr_success": False,
        "notes": "",
    }

    if not DEPS["pymupdf"]:
        result["document_type"] = "UNSUPPORTED"
        result["notes"] = "pymupdf not installed; cannot analyze PDF."
        return result

    doc = pymupdf.open(path)
    try:
        result["page_count"] = doc.page_count

        meta_title = ""
        try:
            meta = doc.metadata or {}
            meta_title = clean_text(meta.get("title", "") or "")
        except Exception:
            pass

        pages_to_scan = min(doc.page_count, MAX_PAGES_TO_SCAN)
        page_texts = []
        chars_per_page = []
        for i in range(pages_to_scan):
            try:
                page = doc.load_page(i)
                ptext = page.get_text("text") or ""
            except Exception:
                ptext = ""
            ptext = clean_text(ptext)
            page_texts.append(ptext)
            chars_per_page.append(len(ptext))

        total_chars = sum(chars_per_page)
        avg_chars = (total_chars / len(chars_per_page)) if chars_per_page else 0

        first_page_text = " ".join(page_texts[:FIRST_PAGES_COUNT])
        body_text = " ".join(page_texts)[:MAX_BODY_CHARS]

        result["first_page_text"] = first_page_text
        result["body_text"] = body_text
        result["text_chars"] = total_chars
        result["text_words"] = len(body_text.split())
        result["headings"] = guess_headings_from_text(
            " \n".join(page_texts[:FIRST_PAGES_COUNT])
        )

        # Title detection: metadata title, else first meaningful line of page 1
        if meta_title:
            result["title_detected"] = meta_title
        else:
            first_lines = [l.strip() for l in page_texts[0].splitlines() if l.strip()] if page_texts else []
            result["title_detected"] = first_lines[0][:150] if first_lines else ""

        # Decide if likely scanned
        looks_scanned = avg_chars < MIN_CHARS_PER_PAGE_NOT_SCANNED
        if looks_scanned:
            result["document_type"] = "SCANNED_PDF"

            if OCR_AVAILABLE:
                result["ocr_attempted"] = True
                ocr_texts = []
                pages_for_ocr = min(doc.page_count, MAX_OCR_PAGES)
                try:
                    for i in range(pages_for_ocr):
                        page = doc.load_page(i)
                        pix = page.get_pixmap(dpi=200)
                        img_bytes = pix.tobytes("png")
                        img = Image.open(io.BytesIO(img_bytes))
                        ocr_text = pytesseract.image_to_string(img)
                        ocr_texts.append(clean_text(ocr_text))
                    ocr_full = " ".join(ocr_texts)
                    if len(ocr_full) > 20:
                        result["ocr_success"] = True
                        result["first_page_text"] = (result["first_page_text"] + " " + ocr_texts[0]).strip() if ocr_texts else result["first_page_text"]
                        result["body_text"] = (result["body_text"] + " " + ocr_full)[:MAX_BODY_CHARS]
                        result["text_chars"] += len(ocr_full)
                        result["text_words"] = len(result["body_text"].split())
                        result["headings"] = result["headings"] or guess_headings_from_text(ocr_texts[0] if ocr_texts else "")
                        result["notes"] = "Scanned PDF; OCR succeeded on first pages."
                    else:
                        result["ocr_success"] = False
                        result["notes"] = "Scanned PDF; OCR attempted but produced little/no text."
                except Exception as e:
                    result["ocr_success"] = False
                    result["notes"] = f"Scanned PDF; OCR attempt failed: {e}"
            else:
                result["notes"] = "Scanned PDF; OCR_NOT_AVAILABLE (missing pymupdf/pytesseract/Pillow)."
        else:
            result["notes"] = "Text extracted normally."

    finally:
        doc.close()

    return result


def extract_docx(path):
    result = {
        "page_count": None, "text_chars": 0, "text_words": 0,
        "first_page_text": "", "body_text": "", "headings": [],
        "title_detected": "", "document_type": "DOCX",
        "ocr_attempted": False, "ocr_success": False, "notes": "",
    }
    if not DEPS["docx"]:
        result["document_type"] = "UNSUPPORTED"
        result["notes"] = "python-docx not installed; cannot analyze DOCX."
        return result

    d = docx.Document(str(path))
    paragraphs = [clean_text(p.text) for p in d.paragraphs if clean_text(p.text)]
    full_text = "\n".join(paragraphs)

    headings = []
    for p in d.paragraphs:
        style_name = (p.style.name if p.style else "") or ""
        txt = clean_text(p.text)
        if txt and "heading" in style_name.lower():
            headings.append(txt)
    if not headings:
        headings = guess_headings_from_text(full_text)

    try:
        core_title = clean_text(d.core_properties.title or "")
    except Exception:
        core_title = ""
    title_detected = core_title or (paragraphs[0][:150] if paragraphs else "")

    body_text = full_text[:MAX_BODY_CHARS]
    result.update({
        "text_chars": len(full_text),
        "text_words": len(full_text.split()),
        "first_page_text": " ".join(paragraphs[:15]),
        "body_text": body_text,
        "headings": headings[:25],
        "title_detected": title_detected,
        "notes": "DOCX parsed via python-docx.",
    })
    return result


def extract_xlsx(path):
    result = {
        "page_count": None, "text_chars": 0, "text_words": 0,
        "first_page_text": "", "body_text": "", "headings": [],
        "title_detected": "", "document_type": "XLSX",
        "ocr_attempted": False, "ocr_success": False, "notes": "",
    }
    if not DEPS["openpyxl"]:
        result["document_type"] = "UNSUPPORTED"
        result["notes"] = "openpyxl not installed; cannot analyze XLSX."
        return result

    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as e:
        result["document_type"] = "UNSUPPORTED"
        result["notes"] = f"Could not open workbook: {e}"
        return result

    sheet_names = wb.sheetnames
    texts = []
    row_limit_per_sheet = 200
    for sname in sheet_names[:10]:
        ws = wb[sname]
        texts.append(f"SHEET: {sname}")
        row_count = 0
        for row in ws.iter_rows(values_only=True):
            if row_count >= row_limit_per_sheet:
                break
            cells = [str(c) for c in row if c is not None]
            if cells:
                texts.append(" ".join(cells))
            row_count += 1

    full_text = clean_text("\n".join(texts))
    body_text = full_text[:MAX_BODY_CHARS]
    result.update({
        "text_chars": len(full_text),
        "text_words": len(full_text.split()),
        "first_page_text": full_text[:2000],
        "body_text": body_text,
        "headings": sheet_names[:25],
        "title_detected": sheet_names[0] if sheet_names else "",
        "notes": f"XLSX parsed; {len(sheet_names)} sheet(s) scanned (partial rows).",
    })
    return result


def extract_pptx(path):
    result = {
        "page_count": None, "text_chars": 0, "text_words": 0,
        "first_page_text": "", "body_text": "", "headings": [],
        "title_detected": "", "document_type": "PPTX",
        "ocr_attempted": False, "ocr_success": False, "notes": "",
    }
    if not DEPS["pptx"]:
        result["document_type"] = "UNSUPPORTED"
        result["notes"] = "python-pptx not installed; cannot analyze PPTX."
        return result

    prs = Presentation(str(path))
    slide_texts = []
    headings = []
    for i, slide in enumerate(prs.slides):
        slide_lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    txt = clean_text("".join(run.text for run in para.runs))
                    if txt:
                        slide_lines.append(txt)
        if slide_lines:
            if i == 0:
                headings.append(slide_lines[0])
            slide_texts.append(" ".join(slide_lines))

    full_text = "\n".join(slide_texts)
    body_text = full_text[:MAX_BODY_CHARS]
    result.update({
        "page_count": len(prs.slides._sldIdLst),
        "text_chars": len(full_text),
        "text_words": len(full_text.split()),
        "first_page_text": " ".join(slide_texts[:2]),
        "body_text": body_text,
        "headings": (headings + guess_headings_from_text(full_text))[:25],
        "title_detected": slide_texts[0][:150] if slide_texts else "",
        "notes": "PPTX parsed via python-pptx.",
    })
    return result


def extract_plain_text_file(path, doc_type):
    result = {
        "page_count": None, "text_chars": 0, "text_words": 0,
        "first_page_text": "", "body_text": "", "headings": [],
        "title_detected": "", "document_type": doc_type,
        "ocr_attempted": False, "ocr_success": False, "notes": "",
    }
    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            raw = f.read(MAX_BODY_CHARS * 2)
    except Exception as e:
        result["document_type"] = "UNSUPPORTED"
        result["notes"] = f"Could not read file: {e}"
        return result

    full_text = clean_text(raw)
    body_text = full_text[:MAX_BODY_CHARS]
    lines = [l for l in full_text.splitlines() if l.strip()]
    result.update({
        "text_chars": len(full_text),
        "text_words": len(full_text.split()),
        "first_page_text": " ".join(lines[:20]),
        "body_text": body_text,
        "headings": guess_headings_from_text(full_text),
        "title_detected": lines[0][:150] if lines else "",
        "notes": f"{doc_type} read as plain text.",
    })
    return result


def extract_csv(path):
    result = extract_plain_text_file(path, "CSV")
    return result


def extract_json_file(path):
    result = {
        "page_count": None, "text_chars": 0, "text_words": 0,
        "first_page_text": "", "body_text": "", "headings": [],
        "title_detected": "", "document_type": "JSON",
        "ocr_attempted": False, "ocr_success": False, "notes": "",
    }
    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            raw = f.read(MAX_BODY_CHARS * 2)
        try:
            data = json.loads(raw)
            flat = json.dumps(data, ensure_ascii=False)
        except Exception:
            flat = raw
    except Exception as e:
        result["document_type"] = "UNSUPPORTED"
        result["notes"] = f"Could not read JSON: {e}"
        return result

    full_text = clean_text(flat)
    body_text = full_text[:MAX_BODY_CHARS]
    result.update({
        "text_chars": len(full_text),
        "text_words": len(full_text.split()),
        "first_page_text": full_text[:2000],
        "body_text": body_text,
        "headings": [],
        "title_detected": "",
        "notes": "JSON parsed and flattened to text.",
    })
    return result


def extract_html(path):
    result = extract_plain_text_file(path, "HTML")
    raw_text = result.get("body_text", "")
    # crude tag strip for better keyword matching (kept simple; no external deps)
    text_no_tags = re.sub(r"<[^>]+>", " ", raw_text)
    text_no_tags = clean_text(text_no_tags)
    title_match = re.search(r"<title>(.*?)</title>", raw_text, re.IGNORECASE | re.DOTALL)
    if title_match:
        result["title_detected"] = clean_text(title_match.group(1))[:150]
    result["body_text"] = text_no_tags[:MAX_BODY_CHARS]
    result["first_page_text"] = text_no_tags[:2000]
    result["headings"] = guess_headings_from_text(text_no_tags)
    result["notes"] = "HTML tags stripped for text analysis."
    return result


EXTRACTORS = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".xlsx": extract_xlsx,
    ".xls": extract_xlsx,  # openpyxl cannot read legacy .xls; will fail gracefully
    ".pptx": extract_pptx,
    ".ppt": extract_pptx,  # python-pptx cannot read legacy .ppt; will fail gracefully
    ".txt": lambda p: extract_plain_text_file(p, "TXT"),
    ".csv": extract_csv,
    ".json": extract_json_file,
    ".html": extract_html,
    ".htm": extract_html,
}


# ==============================================================================
# SCORING
# ==============================================================================

def count_occurrences(text, phrase):
    if not text or not phrase:
        return 0
    return text.count(phrase)


def score_category(category, filename_lower, title_lower, first_page_lower,
                    headings_lower_joined, body_lower):
    groups = KEYWORD_GROUPS[category]
    strong_kws = groups["strong"]
    moderate_kws = groups["moderate"]

    score = 0
    evidence = []

    sections = {
        "filename": filename_lower,
        "title": title_lower,
        "first_page": first_page_lower,
        "headings": headings_lower_joined,
    }

    for section_name, section_text in sections.items():
        weights = SECTION_WEIGHTS[section_name]
        for kw in strong_kws:
            if kw in section_text:
                score += weights["strong"]
                evidence.append(f"{section_name}:{kw}")
        for kw in moderate_kws:
            if kw in section_text:
                score += weights["moderate"]
                evidence.append(f"{section_name}:{kw}")

    # Body frequency scoring, capped per category, with dampening for
    # generic words so they cannot dominate classification.
    body_score = 0
    for kw in strong_kws + moderate_kws:
        occurrences = count_occurrences(body_lower, kw)
        if occurrences <= 0:
            continue
        is_generic = kw in GENERIC_WORDS
        contribution = min(occurrences, 10)  # +1 per occurrence, soft cap
        if is_generic:
            contribution = min(contribution, GENERIC_WORD_CAP)
        body_score += contribution
        if occurrences > 0:
            evidence.append(f"body:{kw}(x{occurrences})")

    body_score = min(body_score, BODY_FREQUENCY_CAP_PER_CATEGORY)
    score += body_score

    # Exact multi-word strong phrase bonus if present verbatim in body
    for kw in strong_kws:
        if " " in kw and kw in body_lower:
            score += EXACT_PHRASE_BONUS
            evidence.append(f"exact_phrase:{kw}")

    # Deduplicate evidence, keep order, cap length for report readability
    seen = set()
    deduped_evidence = []
    for e in evidence:
        if e not in seen:
            seen.add(e)
            deduped_evidence.append(e)
    deduped_evidence = deduped_evidence[:20]

    return score, deduped_evidence


def classify_document(filename, title, first_page_text, headings, body_text):
    filename_lower = filename.lower()
    title_lower = (title or "").lower()
    first_page_lower = (first_page_text or "").lower()
    headings_lower_joined = " ".join(headings).lower() if headings else ""
    body_lower = (body_text or "").lower()

    scores = {}
    evidence_map = {}
    for cat in CATEGORIES:
        s, ev = score_category(
            cat, filename_lower, title_lower, first_page_lower,
            headings_lower_joined, body_lower,
        )
        scores[cat] = s
        evidence_map[cat] = ev

    ranked = sorted(scores.items(), key=lambda kv: kv[1], reverse=True)
    top_cat, top_score = ranked[0]
    second_cat, second_score = ranked[1] if len(ranked) > 1 else (None, 0)

    has_any_content = bool(filename_lower or title_lower or first_page_lower or body_lower)

    if top_score <= 0 or not has_any_content:
        primary_category = "08_Other"
        secondary_categories = []
        confidence = "UNKNOWN"
        combined_evidence = []
    else:
        primary_category = top_cat
        combined_evidence = evidence_map[top_cat]

        secondary_categories = []
        for cat, sc in ranked[1:]:
            if sc >= max(10, top_score * 0.4):
                secondary_categories.append(cat)
        secondary_categories = secondary_categories[:3]

        margin = top_score - second_score
        independent_signal_sections = len(
            {e.split(":", 1)[0] for e in combined_evidence}
        )

        if top_score >= 25 and margin >= 10 and independent_signal_sections >= 2:
            confidence = "HIGH"
        elif top_score >= 12:
            confidence = "MEDIUM"
        elif top_score > 0:
            confidence = "LOW"
        else:
            confidence = "UNKNOWN"

    return {
        "primary_category": primary_category,
        "secondary_categories": secondary_categories,
        "confidence": confidence,
        "scores": scores,
        "evidence": combined_evidence,
    }


# ==============================================================================
# DUPLICATE DETECTION
# ==============================================================================

def normalize_filename_stem(name):
    stem = Path(name).stem.lower()
    stem = re.sub(r"[_\-]?[0-9a-f]{6,}$", "", stem)  # strip trailing hash-like suffixes
    stem = re.sub(r"[^a-z0-9]+", " ", stem).strip()
    return stem


def detect_duplicates(records):
    """
    records: list of dicts with 'filename', 'sha256', 'size_bytes'
    Returns: dict filename -> (duplicate_group_id, duplicate_type, similarity)
    """
    dup_info = {}
    group_counter = 0

    # 1. Exact duplicates via SHA-256
    hash_groups = defaultdict(list)
    for r in records:
        if r.get("sha256"):
            hash_groups[r["sha256"]].append(r["filename"])

    for file_hash, filenames in hash_groups.items():
        if len(filenames) > 1:
            group_counter += 1
            gid = f"EXACT_{group_counter}"
            for fn in filenames:
                dup_info[fn] = {
                    "duplicate_group_id": gid,
                    "duplicate_type": "EXACT_SHA256",
                    "similarity": 1.0,
                }

    # 2. Near-duplicate filename variants (e.g. document.pdf, document_abc123.pdf)
    remaining = [r for r in records if r["filename"] not in dup_info]
    stem_groups = defaultdict(list)
    for r in remaining:
        stem = normalize_filename_stem(r["filename"])
        if stem:
            stem_groups[stem].append(r["filename"])

    for stem, filenames in stem_groups.items():
        if len(filenames) > 1:
            group_counter += 1
            gid = f"NEARNAME_{group_counter}"
            for fn in filenames:
                dup_info[fn] = {
                    "duplicate_group_id": gid,
                    "duplicate_type": "FILENAME_VARIANT",
                    "similarity": None,
                }

    # 3. Near-duplicate content (compare short text fingerprints for files not
    #    already grouped) -- lightweight O(n^2) similarity on truncated text,
    #    only among files of similar size to keep it fast for ~159 files.
    remaining2 = [r for r in records if r["filename"] not in dup_info and r.get("fingerprint")]
    remaining2.sort(key=lambda r: r.get("size_bytes") or 0)
    n = len(remaining2)
    used = set()
    for i in range(n):
        if remaining2[i]["filename"] in used:
            continue
        group_members = [remaining2[i]["filename"]]
        for j in range(i + 1, n):
            if remaining2[j]["filename"] in used:
                continue
            size_i = remaining2[i].get("size_bytes") or 0
            size_j = remaining2[j].get("size_bytes") or 0
            if size_i and size_j:
                ratio = min(size_i, size_j) / max(size_i, size_j)
                if ratio < 0.85:
                    continue
            sim = SequenceMatcher(
                None, remaining2[i]["fingerprint"], remaining2[j]["fingerprint"]
            ).ratio()
            if sim >= 0.9:
                group_members.append(remaining2[j]["filename"])
                used.add(remaining2[j]["filename"])
        if len(group_members) > 1:
            group_counter += 1
            gid = f"NEARCONTENT_{group_counter}"
            for fn in group_members:
                used.add(fn)
                dup_info[fn] = {
                    "duplicate_group_id": gid,
                    "duplicate_type": "NEAR_DUPLICATE_CONTENT",
                    "similarity": round(sim, 3) if fn != remaining2[i]["filename"] else None,
                }

    return dup_info


# ==============================================================================
# MAIN PROCESSING
# ==============================================================================

def analyze_file(path):
    filename = path.name
    ext = path.suffix.lower()
    size_bytes = path.stat().st_size

    record = {
        "filename": filename,
        "full_path": str(path),
        "extension": ext,
        "size_bytes": size_bytes,
        "document_type": "UNKNOWN",
        "page_count": None,
        "text_chars": 0,
        "text_words": 0,
        "ocr_attempted": False,
        "ocr_success": False,
        "primary_category": "08_Other",
        "secondary_categories": [],
        "confidence": "UNKNOWN",
        "scores": {c: 0 for c in CATEGORIES},
        "evidence": [],
        "title_detected": "",
        "duplicate_group_id": None,
        "sha256": None,
        "notes": "",
    }

    record["sha256"] = sha256_file(path)

    if ext not in SUPPORTED_EXTENSIONS:
        record["document_type"] = "UNSUPPORTED"
        record["confidence"] = "UNKNOWN"
        record["notes"] = f"Unsupported file extension: {ext}"
        return record

    extractor = EXTRACTORS.get(ext)
    if extractor is None:
        record["document_type"] = "UNSUPPORTED"
        record["confidence"] = "UNKNOWN"
        record["notes"] = f"No extractor registered for extension: {ext}"
        return record

    try:
        extraction = extractor(path)
    except Exception as e:
        record["document_type"] = "UNKNOWN"
        record["confidence"] = "UNKNOWN"
        record["notes"] = f"Analysis error: {e}"
        return record

    record["document_type"] = extraction.get("document_type", "UNKNOWN")
    record["page_count"] = extraction.get("page_count")
    record["text_chars"] = extraction.get("text_chars", 0)
    record["text_words"] = extraction.get("text_words", 0)
    record["ocr_attempted"] = extraction.get("ocr_attempted", False)
    record["ocr_success"] = extraction.get("ocr_success", False)
    record["title_detected"] = extraction.get("title_detected", "")
    record["notes"] = extraction.get("notes", "")

    if record["document_type"] == "UNSUPPORTED":
        record["confidence"] = "UNKNOWN"
        return record

    try:
        classification = classify_document(
            filename=filename,
            title=extraction.get("title_detected", ""),
            first_page_text=extraction.get("first_page_text", ""),
            headings=extraction.get("headings", []),
            body_text=extraction.get("body_text", ""),
        )
        record["primary_category"] = classification["primary_category"]
        record["secondary_categories"] = classification["secondary_categories"]
        record["confidence"] = classification["confidence"]
        record["scores"] = classification["scores"]
        record["evidence"] = classification["evidence"]
    except Exception as e:
        record["confidence"] = "UNKNOWN"
        record["notes"] = (record["notes"] + f" | Classification error: {e}").strip(" |")

    # Fingerprint for near-duplicate content detection (small, in-memory only)
    body_sample = extraction.get("body_text", "") or ""
    record["_fingerprint"] = body_sample[:3000]

    return record


def main():
    print("=" * 70)
    print("MRPL DEEP KNOWLEDGE BASE CLASSIFIER")
    print("=" * 70)
    print()
    print("SOURCE:")
    print(f"  {SOURCE_DIR}")
    print()
    print("MODE:")
    print("  READ ONLY")
    print()
    print("NO FILES WILL BE:")
    print("  - DELETED")
    print("  - MOVED")
    print("  - RENAMED")
    print("  - MODIFIED")
    print()
    print("=" * 70)
    print()

    print_dependency_report()

    if not SOURCE_DIR.exists():
        print(f"ERROR: Source folder does not exist: {SOURCE_DIR}")
        sys.exit(1)

    all_files = sorted([p for p in SOURCE_DIR.iterdir() if p.is_file()])
    total = len(all_files)
    print(f"Found {total} file(s) in 08_Other.\n")

    records = []
    for idx, path in enumerate(all_files, start=1):
        print(f"[{idx:03d}/{total:03d}] {path.name}")
        try:
            record = analyze_file(path)
        except Exception as e:
            record = {
                "filename": path.name,
                "full_path": str(path),
                "extension": path.suffix.lower(),
                "size_bytes": path.stat().st_size if path.exists() else 0,
                "document_type": "UNKNOWN",
                "page_count": None,
                "text_chars": 0,
                "text_words": 0,
                "ocr_attempted": False,
                "ocr_success": False,
                "primary_category": "08_Other",
                "secondary_categories": [],
                "confidence": "UNKNOWN",
                "scores": {c: 0 for c in CATEGORIES},
                "evidence": [],
                "title_detected": "",
                "duplicate_group_id": None,
                "sha256": None,
                "notes": f"Analysis error: {e}\n{traceback.format_exc(limit=2)}",
                "_fingerprint": "",
            }
        records.append(record)

    print()
    print("Running duplicate detection...")
    dup_input = [
        {
            "filename": r["filename"],
            "sha256": r.get("sha256"),
            "size_bytes": r.get("size_bytes"),
            "fingerprint": r.get("_fingerprint", ""),
        }
        for r in records
    ]
    dup_info = detect_duplicates(dup_input)
    for r in records:
        info = dup_info.get(r["filename"])
        if info:
            r["duplicate_group_id"] = info["duplicate_group_id"]
            r["duplicate_type"] = info["duplicate_type"]
            r["duplicate_similarity"] = info["similarity"]
        else:
            r["duplicate_type"] = None
            r["duplicate_similarity"] = None
        r.pop("_fingerprint", None)

    print("Writing reports...")
    write_json_report(records)
    write_csv_report(records)
    write_summary_report(records)

    confidence_counts = Counter(r["confidence"] for r in records)

    print()
    print("=" * 70)
    print("DEEP CLASSIFICATION COMPLETE")
    print("=" * 70)
    print()
    print(f"FILES ANALYZED: {len(records)}")
    print()
    print(f"HIGH: {confidence_counts.get('HIGH', 0)}")
    print(f"MEDIUM: {confidence_counts.get('MEDIUM', 0)}")
    print(f"LOW: {confidence_counts.get('LOW', 0)}")
    print(f"UNKNOWN: {confidence_counts.get('UNKNOWN', 0)}")
    print()
    print("REPORTS:")
    print(f"  {JSON_REPORT_PATH}")
    print(f"  {CSV_REPORT_PATH}")
    print(f"  {SUMMARY_REPORT_PATH}")
    print()
    print("NO SOURCE FILES WERE MODIFIED.")
    print("NO FILES WERE MOVED.")
    print("NO FILES WERE DELETED.")
    print("=" * 70)


# ==============================================================================
# REPORT WRITERS
# ==============================================================================

def write_json_report(records):
    output = []
    for r in records:
        output.append({
            "filename": r["filename"],
            "full_path": r["full_path"],
            "extension": r["extension"],
            "size_bytes": r["size_bytes"],
            "document_type": r["document_type"],
            "page_count": r["page_count"],
            "text_chars": r["text_chars"],
            "text_words": r["text_words"],
            "ocr_attempted": r["ocr_attempted"],
            "ocr_success": r["ocr_success"],
            "primary_category": r["primary_category"],
            "secondary_categories": r["secondary_categories"],
            "confidence": r["confidence"],
            "scores": r["scores"],
            "evidence": r["evidence"],
            "title_detected": r["title_detected"],
            "duplicate_group_id": r["duplicate_group_id"],
            "duplicate_type": r.get("duplicate_type"),
            "duplicate_similarity": r.get("duplicate_similarity"),
            "notes": r["notes"],
        })

    with open(JSON_REPORT_PATH, "w", encoding="utf-8") as f:
        json.dump(
            {
                "generated_at": datetime.now().isoformat(),
                "source_folder": str(SOURCE_DIR),
                "total_files": len(records),
                "files": output,
            },
            f,
            indent=2,
            ensure_ascii=False,
        )


CSV_CATEGORY_COLUMN_MAP = {
    "score_refinery": "01_Refinery_Manufacturing",
    "score_environment": "02_Environment_Compliance",
    "score_safety": "03_Safety_HSE",
    "score_company": "04_Company_General",
    "score_policy": "05_Policies_Certifications",
    "score_website": "06_Website_Content",
    "score_finance": "07_Finance",
}

CSV_COLUMNS = [
    "filename", "full_path", "extension", "size_bytes", "document_type",
    "page_count", "text_chars", "text_words", "ocr_attempted", "ocr_success",
    "primary_category", "secondary_categories", "confidence",
    "score_refinery", "score_environment", "score_safety", "score_company",
    "score_policy", "score_website", "score_finance",
    "title_detected", "evidence", "duplicate_group", "notes",
]


def write_csv_report(records):
    with open(CSV_REPORT_PATH, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(CSV_COLUMNS)
        for r in records:
            row = [
                r["filename"],
                r["full_path"],
                r["extension"],
                r["size_bytes"],
                r["document_type"],
                r["page_count"] if r["page_count"] is not None else "",
                r["text_chars"],
                r["text_words"],
                r["ocr_attempted"],
                r["ocr_success"],
                r["primary_category"],
                ";".join(r["secondary_categories"]),
                r["confidence"],
            ]
            for csv_col, cat_name in CSV_CATEGORY_COLUMN_MAP.items():
                row.append(r["scores"].get(cat_name, 0))
            row.extend([
                r["title_detected"],
                " | ".join(r["evidence"]),
                r["duplicate_group_id"] or "",
                r["notes"],
            ])
            writer.writerow(row)


def write_summary_report(records):
    total = len(records)

    category_counts = Counter(r["primary_category"] for r in records)
    confidence_counts = Counter(r["confidence"] for r in records)

    doc_type_counts = Counter()
    for r in records:
        dt = r["document_type"]
        if dt == "TEXT_PDF":
            doc_type_counts["Text PDFs"] += 1
        elif dt == "SCANNED_PDF":
            doc_type_counts["Scanned PDFs"] += 1
        elif dt == "TXT":
            doc_type_counts["TXT"] += 1
        elif dt == "DOCX":
            doc_type_counts["DOCX"] += 1
        elif dt in ("XLSX",):
            doc_type_counts["XLSX"] += 1
        elif dt == "PPTX":
            doc_type_counts["PPTX"] += 1
        else:
            doc_type_counts["Other"] += 1

    ocr_attempted = sum(1 for r in records if r["ocr_attempted"])
    ocr_successful = sum(1 for r in records if r["ocr_success"])
    ocr_unavailable = sum(
        1 for r in records
        if r["document_type"] == "SCANNED_PDF" and not r["ocr_attempted"]
    )
    ocr_failed = sum(
        1 for r in records
        if r["ocr_attempted"] and not r["ocr_success"]
    )

    exact_groups = {r["duplicate_group_id"] for r in records if r.get("duplicate_type") == "EXACT_SHA256"}
    near_groups = {
        r["duplicate_group_id"] for r in records
        if r.get("duplicate_type") in ("FILENAME_VARIANT", "NEAR_DUPLICATE_CONTENT")
    }

    lines = []
    lines.append("=" * 70)
    lines.append("MRPL DEEP KNOWLEDGE BASE CLASSIFIER - SUMMARY")
    lines.append("=" * 70)
    lines.append(f"Generated: {datetime.now().isoformat()}")
    lines.append(f"Source folder: {SOURCE_DIR}")
    lines.append("")
    lines.append(f"Total files analyzed: {total}")
    lines.append("")
    lines.append("Category suggestions (primary_category):")
    for cat in ALL_CATEGORIES:
        lines.append(f"  {cat:28s}: {category_counts.get(cat, 0)}")
    lines.append("")
    lines.append("Confidence:")
    lines.append(f"  HIGH   : {confidence_counts.get('HIGH', 0)}")
    lines.append(f"  MEDIUM : {confidence_counts.get('MEDIUM', 0)}")
    lines.append(f"  LOW    : {confidence_counts.get('LOW', 0)}")
    lines.append(f"  UNKNOWN: {confidence_counts.get('UNKNOWN', 0)}")
    lines.append("")
    lines.append("Document types:")
    for label in ["Text PDFs", "Scanned PDFs", "TXT", "DOCX", "XLSX", "PPTX", "Other"]:
        lines.append(f"  {label:14s}: {doc_type_counts.get(label, 0)}")
    lines.append("")
    lines.append("OCR:")
    lines.append(f"  OCR attempted : {ocr_attempted}")
    lines.append(f"  OCR successful: {ocr_successful}")
    lines.append(f"  OCR unavailable/not attempted on scanned docs: {ocr_unavailable}")
    lines.append(f"  OCR failed    : {ocr_failed}")
    lines.append("")
    lines.append("Duplicates:")
    lines.append(f"  Exact duplicate groups (SHA-256): {len(exact_groups)}")
    lines.append(f"  Near-duplicate groups (filename/content): {len(near_groups)}")
    lines.append("")
    lines.append("Reports written:")
    lines.append(f"  {JSON_REPORT_PATH}")
    lines.append(f"  {CSV_REPORT_PATH}")
    lines.append(f"  {SUMMARY_REPORT_PATH}")
    lines.append("")
    lines.append("NO SOURCE FILES WERE MODIFIED, MOVED, OR DELETED.")
    lines.append("=" * 70)

    with open(SUMMARY_REPORT_PATH, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


if __name__ == "__main__":
    main()